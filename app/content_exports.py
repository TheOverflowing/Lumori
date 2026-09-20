"""Portable, in-memory downloads of an existing material version.

This module does not fetch resources or call a model. Exports group citations by
document; the saved asset and chunk-level provenance identifiers stay intact.
"""
from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from io import BytesIO
import json
from pathlib import Path
import re
from threading import Lock
from typing import Any, Mapping
from urllib.parse import urlsplit, urlunsplit
from uuid import uuid4
from xml.sax.saxutils import escape


@dataclass(frozen=True)
class ExportedMaterial:
    data: bytes
    mime_type: str
    filename: str


@dataclass(frozen=True)
class _Block:
    kind: str
    text: str


_CONTEXT = re.compile(r'^context_[a-f0-9]{32}$', re.I)
_FRAGMENTS = re.compile(
    r'(`{3,})[\s\S]*?(?:\1|$)|(`{1,2})[^\n]*?\2|'
    r'\[([^\[\]\n]+)\]|\(([^()\n]+)\)|（([^（）\n]+)）'
)
_INVALID_XML = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')
_CJK = re.compile('[\u2e80-\u9fff\uf900-\ufaff]')
_FONT_DIR = Path(__file__).with_name('fonts')
_FONTS = {
    'LumoriSans': ('NotoSans-Regular.ttf', 'Noto Sans', 'embedRegular'),
    'LumoriSansBold': ('NotoSans-Bold.ttf', 'Noto Sans', 'embedBold'),
    'LumoriCJK': ('NotoSansSC-Regular.ttf', 'Noto Sans SC', 'embedRegular'),
    'LumoriMath': ('NotoSansMath-Regular.ttf', 'Noto Sans Math', 'embedRegular'),
    'LumoriMono': ('NotoSansMono-Regular.ttf', 'Noto Sans Mono', 'embedRegular'),
}
_FONT_LOCK = Lock()


def _text(value: Any) -> str:
    return _INVALID_XML.sub('', '' if value is None else str(value))


@lru_cache(maxsize=1)
def _registered_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    with _FONT_LOCK:
        for name, (filename, _, _) in _FONTS.items():
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, str(_FONT_DIR / filename)))
        return {name: pdfmetrics.getFont(name).face.charToGlyph for name in _FONTS}


def _font_runs(text: str, *, bold=False, code=False):
    """Select an embedded font per run, including mixed Chinese and mathematics.

    Noto SC retains its complete upstream regional glyph set. Unsupported
    characters remain identifiable as a Unicode codepoint instead of disappearing
    or becoming an unexplained square. No host fonts are consulted.
    """
    coverage = _registered_fonts()
    primary = 'LumoriMono' if code else 'LumoriSansBold' if bold else 'LumoriSans'
    candidates = list(dict.fromkeys([primary, 'LumoriSans', 'LumoriMath', 'LumoriCJK']))
    previous = None
    text_run = ''
    for character in text:
        name = next((font for font in candidates if ord(character) in coverage[font]), None)
        if character in '\n\t':
            name = primary
        if name is None:
            name, character = primary, f'[U+{ord(character):04X}]'
        if name != previous and text_run:
            yield previous, text_run
            text_run = ''
        previous, text_run = name, text_run + character
    if text_run:
        yield previous, text_run


def _json(value: Any, fallback: Any) -> Any:
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except (TypeError, ValueError):
            return fallback
    return value if isinstance(value, type(fallback)) else fallback


def _ids(values: Any) -> list[str]:
    return list(dict.fromkeys(value for value in values
        if isinstance(value, str) and value.strip())) if isinstance(values, list) else []


def _external(source: dict) -> dict:
    metadata = source.get('metadata') if isinstance(source.get('metadata'), dict) else {}
    external = source.get('external_source') or metadata.get('external_source')
    return external if isinstance(external, dict) else {}


def _source_url(source: dict) -> str | None:
    external = _external(source)
    return next((value for value in (external.get('reading_url'), external.get('url'))
                 if _safe_url(value)), None)


def _document_key(source: dict, source_id: str) -> tuple[str, str]:
    document_id = source.get('document_id')
    if isinstance(document_id, str) and document_id.strip():
        return ('document', document_id)
    url = _source_url(source)
    if url:
        parsed = urlsplit(url)
        # Anchor targets identify passages in the same web document. Keep path
        # and query exact so different pages, translations or versions stay apart.
        host = parsed.hostname.lower()
        host = '[' + host + ']' if ':' in host else host
        port = parsed.port
        if port and (parsed.scheme, port) not in (('http', 80), ('https', 443)):
            host += ':' + str(port)
        return ('url', urlunsplit((parsed.scheme, host, parsed.path or '/', parsed.query, '')))
    # A shared filename is not reliable document identity.
    return ('chunk', source_id)


def _page_ranges(pages: list[int], zh: bool) -> str:
    ranges = []
    for page in sorted(set(pages)):
        if ranges and ranges[-1][1] + 1 == page:
            ranges[-1][1] = page
        else:
            ranges.append([page, page])
    return ('、' if zh else ', ').join(str(start) if start == end else f'{start}-{end}'
                                      for start, end in ranges)


class _Citations:
    def __init__(self, sources: list, asset: dict, zh: bool, include_answers: bool):
        self.records: dict[str, tuple[tuple[str, str], dict | None]] = {}
        self.used: set[str] = set()
        self.unknown = False
        self.zh = zh
        for source in sources:
            if not isinstance(source, dict):
                continue
            id_ = source.get('id')
            if isinstance(id_, str) and id_ and id_ not in self.records:
                self.records[id_] = (_document_key(source, id_), source)

        def visit(value):
            if isinstance(value, list):
                for child in value:
                    visit(child)
            elif isinstance(value, dict):
                for id_ in _ids(value.get('citation_ids')):
                    if id_ not in self.records:
                        self.records[id_] = (('missing', id_), None)
                for key, child in value.items():
                    if key != 'citation_ids':
                        visit(child)
        visit(asset)

        # Number only documents referenced by fields actually exported, retaining
        # their source-array order. Hidden answer-only references cannot leave gaps.
        planned = set()

        def scan(value):
            for _, ids in self._matches(value):
                planned.update(ids)

        scan(asset.get('title'))
        for objective in asset.get('learning_objectives') or []:
            scan(objective)
        for section in asset.get('sections') or []:
            if isinstance(section, dict):
                scan(section.get('heading'))
                scan(section.get('text'))
                planned.update(_ids(section.get('citation_ids')))
        for question in asset.get('questions') or []:
            if not isinstance(question, dict):
                continue
            scan(question.get('stem'))
            for option in question.get('options') or []:
                scan(option)
            planned.update(_ids(question.get('citation_ids')))
            if include_answers:
                scan(question.get('answer'))
                scan(question.get('explanation'))
        if not asset.get('sections') and not asset.get('questions'):
            scan(asset.get('evidence_note'))
        active = {self.records[id_][0] for id_ in planned if id_ in self.records}
        order = dict.fromkeys(key for key, _ in self.records.values() if key in active)
        self.numbers = {key: number for number, key in enumerate(order, 1)}

    def _matches(self, value):
        for match in _FRAGMENTS.finditer(_text(value)):
            group = next((part for part in match.groups()[2:] if part is not None), None)
            ids = re.split(r'[\s,;，；、]+', group.strip()) if group else []
            if ids and all(id_ in self.records or _CONTEXT.fullmatch(id_) for id_ in ids):
                yield match, list(dict.fromkeys(ids))

    def _number(self, id_: str) -> str:
        self.used.add(id_)
        record = self.records.get(id_)
        if record is None:
            self.unknown = True
            return '?'
        return str(self.numbers.setdefault(record[0], len(self.numbers) + 1))

    def _marker(self, ids: list[str]) -> str:
        numbers = dict.fromkeys(self._number(id_) for id_ in ids)
        return '[' + ', '.join(numbers) + ']'

    def convert(self, value: Any) -> tuple[str, set[str]]:
        content = _text(value)
        inline: set[str] = set()

        parts = []
        position = 0
        for match, ids in self._matches(content):
            parts.append(content[position:match.start()])
            inline.update(ids)
            parts.append(self._marker(ids))
            position = match.end()
        parts.append(content[position:])
        return ''.join(parts), inline

    def references(self, ids: Any, inline: set[str]) -> str:
        declared = _ids(ids)
        # Declared supporting chunks still contribute page locations even when
        # their document already has an inline citation and needs no extra footer.
        for id_ in declared:
            self._number(id_)
        inline_groups = {self.records[id_][0] for id_ in inline if id_ in self.records}
        remaining = [id_ for id_ in declared if id_ not in inline and
                     (id_ not in self.records or self.records[id_][0] not in inline_groups)]
        return ('来源：' if self.zh else 'Sources: ') + self._marker(remaining) if remaining else ''

    def bibliography(self) -> list[str]:
        result = []
        groups = {}
        for id_, (key, source) in self.records.items():
            if id_ in self.used:
                groups.setdefault(key, []).append(source)
        for key in sorted(groups, key=self.numbers.__getitem__):
            number = self.numbers[key]
            sources = [source for source in groups[key] if source is not None]
            if not sources:
                result.append(f'[{number}] ' + ('来源暂不可用。' if self.zh else 'Source unavailable.'))
                continue

            def values(field):
                return list(dict.fromkeys(_text(_external(source).get(field)) for source in sources
                                          if _external(source).get(field)))

            titles = values('title') or [_text(source['document_name']) for source in sources
                                         if source.get('document_name')]
            title = titles[0] if titles else ('未命名资料' if self.zh else 'Untitled material')
            line = f'[{number}] {_text(title)}'
            pages = sorted({source['page'] for source in sources
                            if type(source.get('page')) is int and source['page'] > 0})
            if pages:
                locations = _page_ranges(pages, self.zh)
                line += f' · 第 {locations} 页' if self.zh else f' · {"p." if len(pages) == 1 else "pp."} {locations}'
            details = [line]
            if attributions := values('attribution'):
                details.append('; '.join(attributions))
            if url := next((_source_url(source) for source in sources if _source_url(source)), None):
                details.append(url)
            if licenses := values('license'):
                details.append(('许可：' if self.zh else 'License: ') + '; '.join(licenses))
                details.extend(value for value in values('license_url') if _safe_url(value))
            result.append('\n'.join(details))
        if self.unknown:
            result.append('[?] ' + ('来源暂不可用。' if self.zh else 'Source unavailable.'))
        return result


def _safe_url(value: Any) -> bool:
    if not isinstance(value, str) or any(ord(c) < 32 for c in value):
        return False
    try:
        url = urlsplit(value)
        url.port  # Reject invalid port syntax before grouping or exporting links.
        return url.scheme in ('http', 'https') and bool(url.hostname) and not url.username and not url.password
    except ValueError:
        return False


def _paragraphs(text: str, kind='body') -> list[_Block]:
    # Fenced code is still authored content. Preserve its text and indentation,
    # while allowing long code lines and paragraphs to wrap at the page edge.
    blocks = []
    start = 0
    for match in re.finditer(r'(?m)^(`{3,})[^\n]*\n([\s\S]*?)(?:^\1[ \t]*$|\Z)', text):
        blocks.extend(_Block(kind, part) for part in re.split(r'\n[ \t]*\n', text[start:match.start()]) if part.strip())
        blocks.append(_Block('code', match[2].rstrip('\n')))
        start = match.end()
    blocks.extend(_Block(kind, part) for part in re.split(r'\n[ \t]*\n', text[start:]) if part.strip())
    return blocks


def _material_blocks(row: Mapping[str, Any], include_answers: bool, language: str):
    asset = _json(row.get('asset'), {})
    sources = _json(row.get('sources'), [])
    zh = language == 'zh'
    citations = _Citations(sources, asset, zh, include_answers)
    title = citations.convert(_text(asset.get('title')) or ('学习资料' if zh else 'Learning material'))[0]
    version = row.get('version', 1)
    version = version if type(version) is int and version > 0 else 1
    blocks = [_Block('title', title),
              _Block('meta', f'Lumori · {"版本" if zh else "Version"} {version}')]

    def add(value, kind='body'):
        converted, inline = citations.convert(value)
        blocks.extend(_paragraphs(converted, kind))
        return inline

    objectives = asset.get('learning_objectives') or []
    if objectives:
        blocks.append(_Block('heading', '学习目标' if zh else 'Learning objectives'))
        for index, objective in enumerate(objectives, 1):
            add(f'{index}. {_text(objective)}')
    for section in asset.get('sections') or []:
        if not isinstance(section, dict):
            continue
        inline = add(section.get('heading'), 'heading')
        inline.update(add(section.get('text')))
        refs = citations.references(section.get('citation_ids'), inline)
        if refs:
            blocks.append(_Block('reference', refs))
    questions = [q for q in asset.get('questions') or [] if isinstance(q, dict)]
    if questions:
        blocks.append(_Block('heading', '练习题' if zh else 'Practice questions'))
    for index, question in enumerate(questions, 1):
        inline = add(f'{index}. {_text(question.get("stem"))}', 'question')
        for option_index, option in enumerate(question.get('options') or []):
            inline.update(add(f'{chr(65 + option_index)}. {_text(option)}', 'option'))
        refs = citations.references(question.get('citation_ids'), inline)
        if refs:
            blocks.append(_Block('reference', refs))
    if include_answers and any(q.get('answer') is not None or q.get('explanation') for q in questions):
        blocks.append(_Block('answers_heading', '参考答案与解析' if zh else 'Answers and explanations'))
        for index, question in enumerate(questions, 1):
            add(f'{index}. {_text(question.get("answer"))}', 'answer')
            add(question.get('explanation'))
    if not asset.get('sections') and not questions and asset.get('evidence_note'):
        add(asset['evidence_note'])
    bibliography = citations.bibliography()
    if bibliography:
        blocks.append(_Block('heading', '参考来源' if zh else 'References'))
        blocks.extend(_Block('bibliography', text) for text in bibliography)
    return title, version, blocks


def _filename(title: str, version: int, format: str) -> str:
    stem = re.sub(r'[\x00-\x1f\x7f/\\:*?"<>|]', '-', title).strip(' .-')
    # Keep the filename comfortably below common 255-byte filesystem limits.
    stem = stem.encode('utf-8')[:150].decode('utf-8', errors='ignore').strip(' .-') or 'Lumori'
    return f'{stem}-v{version}.{format}'


def _pdf(blocks: list[_Block], title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

    _registered_fonts()
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=A4, rightMargin=23 * mm, leftMargin=23 * mm,
        topMargin=22 * mm, bottomMargin=22 * mm, title=title, author='Lumori',
        allowSplitting=1, pageCompression=1)
    settings = {
        'title': dict(fontSize=23, leading=31, spaceAfter=8, keepWithNext=True),
        'meta': dict(fontSize=9, leading=14, spaceAfter=20, textColor=colors.HexColor('#606963'), keepWithNext=True),
        'heading': dict(fontSize=14, leading=21, spaceBefore=15, spaceAfter=8, keepWithNext=True),
        'answers_heading': dict(fontSize=14, leading=21, spaceAfter=8, keepWithNext=True),
        'body': dict(fontSize=10.5, leading=17, spaceAfter=9),
        'question': dict(fontSize=11, leading=18, spaceBefore=11, spaceAfter=7, keepWithNext=True),
        'answer': dict(fontSize=11, leading=18, spaceBefore=11, spaceAfter=7, keepWithNext=True),
        'option': dict(fontSize=10.5, leading=17, leftIndent=12, spaceAfter=4),
        'reference': dict(fontSize=8.5, leading=13, spaceAfter=8, textColor=colors.HexColor('#606963')),
        'bibliography': dict(fontSize=9, leading=14, spaceAfter=12),
        'code': dict(fontSize=9, leading=14, spaceBefore=4, spaceAfter=10, leftIndent=9,
                     rightIndent=9, backColor=colors.HexColor('#f3f5f3'), borderPadding=7),
    }
    story = []
    for block in blocks:
        if block.kind == 'answers_heading':
            story.append(PageBreak())
        bold = block.kind in ('title', 'heading', 'answers_heading', 'question', 'answer')
        style = ParagraphStyle(block.kind, fontName='LumoriSans', alignment=TA_LEFT,
            wordWrap='CJK' if _CJK.search(block.text) else None,
            splitLongWords=True, allowWidows=0, allowOrphans=0, **settings[block.kind])
        parts = []
        for font, run in _font_runs(block.text, bold=bold, code=block.kind == 'code'):
            text = escape(run).replace('\n', '<br/>')
            if block.kind == 'code':
                text = text.replace('  ', '&#160;&#160;').replace('\t', '&#160;' * 4)
            parts.append(f'<font name="{font}">{text}</font>')
        text = ''.join(parts)
        story.append(Paragraph(text, style))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont('LumoriSans', 8)
        canvas.setFillColor(colors.HexColor('#747b76'))
        canvas.drawRightString(A4[0] - 23 * mm, 12 * mm, str(document.page))
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()


def _docx(blocks: list[_Block], title: str, language: str) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(22), Mm(22)
    section.left_margin, section.right_margin = Mm(23), Mm(23)
    doc.core_properties.title = title
    doc.core_properties.author = 'Lumori'
    doc.core_properties.subject = ''
    doc.core_properties.comments = ''
    for name, size in [('Normal', 10.5), ('Title', 23), ('Heading 1', 14), ('Heading 2', 11)]:
        style = doc.styles[name]
        style.font.name = 'Noto Sans'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        props = style.element.get_or_add_rPr()
        fonts = props.find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            props.append(fonts)
        fonts.set(qn('w:eastAsia'), 'Noto Sans SC')
        for attribute in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn('w:' + attribute), None)
        lang = OxmlElement('w:lang')
        lang.set(qn('w:val'), 'zh-CN' if language == 'zh' else 'en-US')
        lang.set(qn('w:eastAsia'), 'zh-CN')
        props.append(lang)
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing = Pt(17)
    normal.space_after = Pt(8)
    normal.widow_control = True
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(16)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(8)
    doc.styles['Heading 1'].paragraph_format.line_spacing = Pt(21)
    doc.styles['Title'].paragraph_format.line_spacing = Pt(31)
    doc.styles['Title'].paragraph_format.space_after = Pt(8)
    # Third-party python-docx runtimes may ship a styled default template.
    # The exported material must not inherit its colored title rule.
    for style in doc.styles:
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    used_fonts = {'LumoriSans', 'LumoriSansBold'}
    for block in blocks:
        style = {'title': 'Title', 'heading': 'Heading 1', 'answers_heading': 'Heading 1'}.get(block.kind, 'Normal')
        paragraph = doc.add_paragraph(style=style)
        if block.kind == 'answers_heading':
            paragraph.paragraph_format.page_break_before = True
        for font, text in _font_runs(block.text, bold=block.kind in ('title', 'heading', 'answers_heading', 'question', 'answer'), code=block.kind == 'code'):
            used_fonts.add(font)
            run = paragraph.add_run(text)
            run.font.name = _FONTS[font][1]
            fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
            for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                fonts.set(qn('w:' + attr), _FONTS[font][1])
            run.bold = font == 'LumoriSansBold'
            if block.kind in ('meta', 'reference', 'bibliography'):
                run.font.size = Pt(9)
            if block.kind in ('meta', 'reference'):
                run.font.color.rgb = RGBColor.from_string('606963')
            if block.kind == 'code':
                run.font.size = Pt(9)
        if block.kind in ('question', 'answer'):
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.keep_with_next = True
        if block.kind == 'meta':
            paragraph.paragraph_format.space_after = Pt(18)
            paragraph.paragraph_format.keep_with_next = True
        if block.kind == 'option':
            paragraph.paragraph_format.left_indent = Mm(4)
            paragraph.paragraph_format.space_after = Pt(4)
        if block.kind == 'bibliography':
            paragraph.paragraph_format.keep_together = True
        if block.kind == 'code':
            paragraph.paragraph_format.left_indent = Mm(3)
            paragraph.paragraph_format.line_spacing = 1.2
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'F3F5F3')
            paragraph._p.get_or_add_pPr().append(shading)
    # A page number is the only recurring footer; content remains editable.
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    run = footer.add_run()
    run.font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    run._r.addnext(field)
    _embed_docx_fonts(doc, used_fonts)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _embed_docx_fonts(doc, used_fonts: set[str]):
    """Package complete editable OFL fonts using OOXML font obfuscation.

    A Word download must remain readable on a machine without Chinese fonts.
    Unlike PDF's automatic glyph subsets, DOCX retains each required full font
    so the recipient can edit the material and type additional characters.
    """
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import qn
    from lxml import etree

    table_part = doc.part.part_related_by(RT.FONT_TABLE)
    table = parse_xml(table_part.blob)
    for index, name in enumerate(sorted(used_fonts), 1):
        filename, family, kind = _FONTS[name]
        key = uuid4()
        # ISO/IEC 29500 stores the GUID bytes in reverse order for the XOR key.
        key_bytes = key.bytes[::-1]
        data = bytearray((_FONT_DIR / filename).read_bytes())
        for offset in range(32):
            data[offset] ^= key_bytes[offset % 16]
        part = Part(PackURI(f'/word/fonts/lumori-{index}.odttf'),
            'application/vnd.openxmlformats-officedocument.obfuscatedFont', bytes(data), doc.part.package)
        relation = table_part.relate_to(part, RT.FONT)
        font = next((entry for entry in table if entry.get(qn('w:name')) == family), None)
        if font is None:
            font = OxmlElement('w:font')
            font.set(qn('w:name'), family)
            table.append(font)
        embed = OxmlElement('w:' + kind)
        embed.set(qn('r:id'), relation)
        embed.set(qn('w:fontKey'), '{' + str(key).upper() + '}')
        embed.set(qn('w:subsetted'), '0')
        font.append(embed)
    table_part._blob = etree.tostring(table, xml_declaration=True, encoding='UTF-8', standalone=True)
    settings = doc.settings.element
    for name in ('embedTrueTypeFonts', 'embedSystemFonts'):
        option = OxmlElement('w:' + name)
        option.set(qn('w:val'), 'true')
        settings.append(option)


def export_material(row: Mapping[str, Any], format: str = 'pdf', include_answers: bool = True,
                    language: str = 'en') -> ExportedMaterial:
    """Export the passed version; authorization/version selection belongs to the route.

    ``asset`` and ``sources`` accept either JSON strings from SQLite or parsed
    objects from the content endpoint. Labels follow ``language``; the authored
    text is never translated or regenerated. PDF and DOCX are generated in memory.
    """
    if format not in ('pdf', 'docx'):
        raise ValueError('Unsupported document format')
    if language not in ('en', 'zh'):
        raise ValueError('Unsupported document language')
    title, version, blocks = _material_blocks(dict(row), include_answers, language)
    data = _pdf(blocks, title) if format == 'pdf' else _docx(blocks, title, language)
    mime = 'application/pdf' if format == 'pdf' else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    return ExportedMaterial(data, mime, _filename(title, version, format))
