"""Portable document exports preserve saved content without provider calls."""
from copy import deepcopy
from io import BytesIO
import json
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader
import pytest

from app.content_exports import export_material, _material_blocks


FIRST = 'context_' + '1' * 32
SECOND = 'context_' + '2' * 32
UNKNOWN = 'context_' + 'f' * 32


@pytest.fixture
def row():
    return {
        'version': 3,
        'asset': {
            'title': 'Machine Learning 机器学习',
            'learning_objectives': ['Explain learning from data. 理解泛化。'],
            'sections': [{
                'heading': 'Learning and generalization',
                'text': f'Training improves a model [{SECOND}].\n\nUse held-out data （{FIRST}；{SECOND}）.',
                'citation_ids': [FIRST, SECOND],
            }],
            'questions': [{
                'stem': 'Which split estimates generalization?',
                'options': ['Training set', 'Test set', 'Both', 'Neither'],
                'answer': 'B', 'explanation': 'Unique hidden explanation: evaluate on unseen data.',
                'citation_ids': [FIRST],
            }],
        },
        'sources': [
            {'id': FIRST, 'document_name': '课程资料.pdf', 'page': 6,
             'external_source': {'title': 'Introduction to machine learning',
                'reading_url': 'https://example.edu/course/chapter',
                'url': 'https://example.edu/raw', 'attribution': 'Course authors',
                'license': 'CC BY 4.0', 'license_url': 'https://creativecommons.org/licenses/by/4.0/'}},
            {'id': SECOND, 'document_name': 'Lecture 2.pdf', 'page': 12},
        ],
        'config': {'language': 'en', 'request_key': 'private-not-for-export'},
    }


def text_of(export):
    if export.filename.endswith('.pdf'):
        return '\n'.join(page.extract_text() for page in PdfReader(BytesIO(export.data)).pages).replace('\u00a0', ' ').replace('\u2060', '')
    return '\n'.join(p.text for p in Document(BytesIO(export.data)).paragraphs).replace('\u00a0', ' ').replace('\u2060', '')


@pytest.mark.parametrize('format', ['pdf', 'docx'])
def test_compatible_six_option_question_exports_every_option_and_its_key(row, format):
    question = row['asset']['questions'][0]
    question['options'].extend(['Fifth supplied option', 'Sixth supplied option'])
    question['answer'] = 'F'
    text = text_of(export_material(row, format, include_answers=True, language='en'))
    assert 'E. Fifth supplied option' in text and 'F. Sixth supplied option' in text
    assert all(option in text for option in question['options'])
    assert '1. F' in text.split('Answers and explanations', 1)[1]


@pytest.mark.parametrize('format', ['pdf', 'docx'])
@pytest.mark.parametrize('language', ['en', 'zh'])
def test_bilingual_export_preserves_content_and_friendly_source_numbers(row, format, language):
    original = deepcopy(row)
    output = export_material(row, format, language=language)
    text = text_of(output)
    assert row == original
    assert 'Machine Learning 机器学习' in text
    assert '理解泛化' in text
    assert 'Training improves a model [2]' in text
    assert 'Use held-out data [1, 2]' in text
    assert 'context_' not in text
    assert 'private-not-for-export' not in text
    assert 'Introduction to machine learning' in text
    assert 'Course authors' in text and 'CC BY 4.0' in text
    assert 'https://example.edu/course/chapter' in text
    assert 'https://example.edu/raw' not in text
    assert 'Unique hidden explanation' in text
    assert output.filename == f'Machine Learning 机器学习-v3.{format}'
    assert output.data.startswith(b'%PDF') if format == 'pdf' else output.data.startswith(b'PK')
    if format == 'docx':
        with ZipFile(BytesIO(output.data)) as archive:
            xml = archive.read('word/document.xml').decode()
            assert 'w:val="Title"' in xml and 'w:val="Heading1"' in xml
            assert 'w:eastAsia="Noto Sans SC"' in archive.read('word/styles.xml').decode()
            assert any(name.startswith('word/fonts/') for name in archive.namelist())
            assert 'embedRegular' in archive.read('word/fontTable.xml').decode()
            assert '<w:pBdr>' not in archive.read('word/styles.xml').decode()
        assert output.mime_type.endswith('wordprocessingml.document')
    else:
        assert output.mime_type == 'application/pdf'


@pytest.mark.parametrize('format', ['pdf', 'docx'])
def test_excluding_answers_omits_answer_heading_and_explanation(row, format):
    output = export_material(row, format, include_answers=False)
    text = text_of(output)
    assert 'Which split estimates generalization?' in text
    assert 'B. Test set' in text
    assert 'Unique hidden explanation' not in text
    assert 'Answers and explanations' not in text
    if format == 'pdf':
        assert len(PdfReader(BytesIO(output.data)).pages) == 1
    else:
        assert not any(p.paragraph_format.page_break_before for p in Document(BytesIO(output.data)).paragraphs)


@pytest.mark.parametrize('format',['pdf','docx'])
@pytest.mark.parametrize('include_answers',[True,False])
def test_questions_only_export_has_no_introduction_and_keeps_answer_choice_independent(row,format,include_answers):
    row['asset'].update(sections=[],learning_objectives=[])
    row['config']['include_explanations']=False
    output=export_material(row,format,include_answers=include_answers)
    text=text_of(output)
    assert 'Which split estimates generalization?' in text
    assert 'B. Test set' in text
    assert 'Learning and generalization' not in text
    assert 'Learning objectives' not in text
    assert ('Unique hidden explanation' in text) is include_answers
    assert 'References' in text and 'context_' not in text


def test_answers_start_on_a_new_page_in_both_formats(row):
    pdf = PdfReader(BytesIO(export_material(row, 'pdf').data))
    assert 'Answers and explanations' not in pdf.pages[0].extract_text()
    assert pdf.pages[1].extract_text().splitlines()[1] == 'Answers and explanations'
    document = Document(BytesIO(export_material(row, 'docx').data))
    heading = next(p for p in document.paragraphs if p.text == 'Answers and explanations')
    assert heading.paragraph_format.page_break_before is True


@pytest.mark.parametrize('format', ['pdf', 'docx'])
def test_json_database_row_and_parsed_endpoint_values_are_equivalent(row, format):
    saved = dict(row)
    for key in ('asset', 'sources', 'config'):
        saved[key] = json.dumps(saved[key], ensure_ascii=False)
    assert text_of(export_material(saved, format)) == text_of(export_material(row, format))


def test_citations_match_frontend_grammar_and_preserve_authored_code(row):
    row['asset']['sections'][0]['text'] = (
        f'First ({SECOND}, {FIRST}, {SECOND}). Second （{FIRST}、{SECOND}）.\n\n'
        f'Keep array[0], [0, 1], (normal notes), [context_window], and `[{FIRST}]`.\n\n'
        f'```python\n[{SECOND}]\n```\n\nMissing [{UNKNOWN}]. Declared [missing-source].'
    )
    row['asset']['sections'][0]['citation_ids'].append('missing-source')
    _, _, blocks = _material_blocks(row, True, 'en')
    text = '\n'.join(block.text for block in blocks).replace('\u00a0', ' ').replace('\u2060', '')
    assert 'First [2, 1]. Second [1, 2].' in text
    assert 'array[0], [0, 1], (normal notes), [context_window]' in text
    assert f'`[{FIRST}]`' in text
    assert any(block.kind == 'code' and block.text == f'[{SECOND}]' for block in blocks)
    assert 'Missing [?]. Declared [3].' in text
    assert '[3] Source unavailable.' in text
    assert '[?] Source unavailable.' in text
    assert UNKNOWN not in text and 'missing-source' not in text


@pytest.mark.parametrize('format', ['pdf', 'docx'])
def test_unsafe_markup_urls_and_invalid_xml_are_text_not_resources(row, format):
    malicious = '<img src="file:///etc/passwd"/><a href="javascript:alert(1)">label</a> & keep'
    row['asset']['sections'][0]['text'] = malicious + '\x00\x08\ud800'
    row['sources'][0]['external_source'].update(reading_url='javascript:alert(1)', url='https://user:password@example.edu/secret')
    row['sources'][0]['external_source']['license_url'] = 'file:///private/data'
    output = export_material(row, format)
    text = text_of(output)
    assert '<img src="file:///etc/passwd"/>' in text
    assert 'javascript:alert(1)' in text  # Authored text is preserved, never fetched.
    assert '& keep' in text
    assert 'user:password' not in text and '/private/data' not in text
    if format == 'pdf':
        assert not any(page.get('/Annots') for page in PdfReader(BytesIO(output.data)).pages)
    else:
        with ZipFile(BytesIO(output.data)) as archive:
            relations = archive.read('word/_rels/document.xml.rels').decode()
            assert 'TargetMode="External"' not in relations


@pytest.mark.parametrize('format', ['pdf', 'docx'])
def test_long_lines_and_multipage_content_do_not_fail_or_drop_text(row, format):
    row['asset']['sections'][0]['text'] = (
        'Start of long content.\n\n' + 'abcdefghij' * 180 + '\n\n' +
        ('这是完整的中文长段落，用于验证分页和换行。' * 160) + '\n\nEnd of long content.'
    )
    output = export_material(row, format)
    text = text_of(output)
    assert 'Start of long content.' in text and 'End of long content.' in text
    assert '完整的中文长段落' in text
    if format == 'pdf':
        assert len(PdfReader(BytesIO(output.data)).pages) >= 3


@pytest.mark.parametrize('format', ['pdf', 'docx'])
def test_missing_optional_fields_and_filename_are_safe(format):
    output = export_material({'asset': {'title': '../../ : 空标题\r\n'}, 'sources': None}, format, language='zh')
    assert '/' not in output.filename and '\\' not in output.filename
    assert '\r' not in output.filename and '\n' not in output.filename
    assert '空标题' in text_of(output)
    empty = export_material({'asset': '{}', 'sources': 'not-json', 'version': None}, format)
    assert empty.filename == f'Learning material-v1.{format}'


def test_unused_and_answer_only_sources_do_not_appear_in_no_answer_download(row):
    row['sources'].append({'id': 'answer-only', 'document_name': 'Answer key document'})
    row['asset']['questions'][0]['explanation'] += ' [answer-only]'
    no_answers = '\n'.join(b.text for b in _material_blocks(row, False, 'en')[2])
    answers = '\n'.join(b.text for b in _material_blocks(row, True, 'en')[2])
    assert 'Answer key document' not in no_answers
    assert '[3] Answer key document' in answers


def test_declared_reference_number_does_not_change_when_inline_source_order_changes(row):
    # Source-order numbering stays the same when a paragraph is edited.
    row['asset']['sections'][0]['text'] = f'Edited [{FIRST}] before [{SECOND}].'
    _, _, blocks = _material_blocks(row, True, 'en')
    assert any(b.text.replace('\u2060', '') == 'Edited [1] before [2].' for b in blocks)
    assert not any(b.kind == 'reference' for b in blocks[:5])


def test_pdf_preserves_greek_and_mathematical_unicode(row):
    formula = 'θ ← θ − η∇L(θ); O(n²), α ≤ β, x ∈ R, ∑ xᵢ.'
    row['asset']['sections'][0]['text'] = formula
    text = text_of(export_material(row, 'pdf'))
    assert formula in text
    assert '■' not in text


def test_pdf_embeds_unicode_fonts_instead_of_relying_on_viewer_language_packs(row):
    output = export_material(row, 'pdf')
    reader = PdfReader(BytesIO(output.data))
    fonts = [font.get_object() for page in reader.pages for font in page['/Resources']['/Font'].values()]
    embedded = [font for font in fonts if '/FontDescriptor' in font]
    assert embedded
    assert all('/FontFile2' in font['/FontDescriptor'] for font in embedded)
    assert not any('STSong' in str(font) for font in fonts)


def test_unsupported_unicode_is_identifiable_instead_of_silently_dropped(row):
    row['asset']['sections'][0]['text'] = 'A rare unsupported glyph: \U0001FAE0.'
    for format in ('pdf', 'docx'):
        assert 'A rare unsupported glyph: [U+1FAE0].' in text_of(export_material(row, format))


@pytest.mark.parametrize('kwargs', [{'format': 'exe'}, {'language': 'es'}])
def test_unsupported_formats_fail_clearly(row, kwargs):
    with pytest.raises(ValueError):
        export_material(row, **kwargs)


@pytest.mark.parametrize('format', ['pdf', 'docx'])
def test_document_chunks_share_one_exported_reference_and_aggregate_pages(row, format):
    row['sources'][0]['document_id'] = 'same-document'
    row['sources'][1]['document_id'] = 'same-document'
    row['sources'][0]['page'], row['sources'][1]['page'] = 1, 2
    row['sources'][1]['external_source'] = deepcopy(row['sources'][0]['external_source'])
    row['asset']['sections'][0]['text'] = f'First [{FIRST}]. Both [{FIRST}, {SECOND}, {FIRST}].'
    row['asset']['questions'][0]['stem'] += f' [{SECOND}]'
    row['asset']['questions'][0]['answer'] += f' [{FIRST}]'
    row['asset']['questions'][0]['explanation'] += f' [{SECOND}]'
    original = deepcopy(row)
    text = text_of(export_material(row, format))
    assert row == original
    assert 'First [1]. Both [1].' in text
    assert 'Which split estimates generalization? [1]' in text
    assert '1. B [1]' in text
    assert 'evaluate on unseen data. [1]' in text
    assert 'pp. 1-2' in text
    assert text.count('[1] Introduction to machine learning') == 1
    assert text.count('https://example.edu/course/chapter') == 1
    assert text.count('Course authors') == 1
    assert text.count('License: CC BY 4.0') == 1
    assert '[2]' not in text and '[1, 1]' not in text and 'context_' not in text
    assert 'Sources:' not in text


def test_inline_document_covers_declared_other_chunk_without_duplicate_footer(row):
    for source in row['sources']:
        source['document_id'] = 'same-document'
    row['asset']['sections'][0]['text'] = f'Text references the document [{FIRST}].'
    row['asset']['sections'][0]['citation_ids'] = [SECOND]
    row['asset']['questions'] = []
    blocks = _material_blocks(row, True, 'en')[2]
    assert not any(b.kind == 'reference' for b in blocks)
    bibliography = [b.text for b in blocks if b.kind == 'bibliography']
    assert len(bibliography) == 1
    assert 'pp. 6, 12' in bibliography[0]


def test_distinct_document_ids_never_merge_even_with_same_filename_and_url(row):
    row['sources'][0]['document_id'] = 'first-version'
    row['sources'][1]['document_id'] = 'second-version'
    for source in row['sources']:
        source['document_name'] = 'same-file.pdf'
        source['external_source'] = {'url': 'https://example.edu/same-file.pdf'}
    blocks = _material_blocks(row, True, 'en')[2]
    bibliography = [b.text for b in blocks if b.kind == 'bibliography']
    assert len(bibliography) == 2
    assert bibliography[0].startswith('[1] same-file.pdf')
    assert bibliography[1].startswith('[2] same-file.pdf')


def test_same_filename_without_document_identity_or_safe_url_does_not_merge(row):
    for source in row['sources']:
        source['document_name'] = 'same-file.pdf'
        source['external_source'] = {'url': 'javascript:alert(1)'}
    blocks = _material_blocks(row, True, 'en')[2]
    assert len([b for b in blocks if b.kind == 'bibliography']) == 2


def test_safe_external_url_fallback_groups_anchor_chunks_and_combines_pages(row):
    row['sources'][0]['external_source']['reading_url'] = 'https://EXAMPLE.edu:443/course/chapter#first'
    row['sources'][1]['metadata'] = {'external_source': {
        'reading_url': 'https://example.edu/course/chapter#second',
        'title': 'Introduction to machine learning', 'license': 'CC BY 4.0',
    }}
    blocks = _material_blocks(row, True, 'en')[2]
    bibliography = [b.text for b in blocks if b.kind == 'bibliography']
    assert len(bibliography) == 1
    assert 'pp. 6, 12' in bibliography[0]
    assert any(b.text.startswith('Training improves a model [1].') for b in blocks)


@pytest.mark.parametrize('second_url', [
    'https://example.edu/course/chapter?language=zh',
    'https://example.edu/course/other',
    'https://example.edu:444/course/chapter',
    'https://user:password@example.edu/course/chapter',
    'https://example.edu:not-a-port/course/chapter',
])
def test_url_fallback_never_collapses_distinct_or_unsafe_addresses(row, second_url):
    row['sources'][1]['external_source'] = {'reading_url': second_url}
    blocks = _material_blocks(row, True, 'en')[2]
    assert len([b for b in blocks if b.kind == 'bibliography']) == 2


def test_explicit_document_identity_does_not_merge_with_anonymous_url_match(row):
    row['sources'][0]['document_id'] = 'known-document'
    row['sources'][1]['external_source'] = deepcopy(row['sources'][0]['external_source'])
    assert len([b for b in _material_blocks(row, True, 'en')[2] if b.kind == 'bibliography']) == 2


def test_only_used_chunk_pages_are_aggregated_and_ranges_are_readable_in_chinese(row):
    for source in row['sources']:
        source['document_id'] = 'same-document'
    row['sources'][0]['page'], row['sources'][1]['page'] = 1, 2
    for id_, page in [('third', 3), ('fifth', 5), ('unused-page', 99)]:
        row['sources'].append({'id': id_, 'document_id': 'same-document', 'page': page,
                               'document_name': '课程资料.pdf'})
    row['asset']['sections'][0]['citation_ids'].extend(['third', 'fifth'])
    bibliography = [b.text for b in _material_blocks(row, True, 'zh')[2] if b.kind == 'bibliography']
    assert len(bibliography) == 1
    assert '第 1-3、5 页' in bibliography[0]
    assert '99' not in bibliography[0]


def test_numbering_compacts_used_documents_when_early_sources_and_answers_are_hidden(row):
    row['sources'].insert(0, {'id': 'unused', 'document_id': 'unused-doc', 'document_name': 'Unused'})
    row['sources'].insert(1, {'id': 'answer-only', 'document_id': 'answers', 'document_name': 'Answer source'})
    row['asset']['questions'][0]['explanation'] += ' [answer-only]'
    without = _material_blocks(row, False, 'en')[2]
    references = [b.text for b in without if b.kind == 'bibliography']
    assert len(references) == 2
    assert references[0].startswith('[1] Introduction') and references[1].startswith('[2] Lecture 2')
    assert not any('Answer source' in b.text or 'Unused' in b.text for b in without)
    assert any('Training improves a model [2]' in b.text for b in without)
    with_answers = _material_blocks(row, True, 'en')[2]
    references = [b.text for b in with_answers if b.kind == 'bibliography']
    assert len(references) == 3
    assert references[0].startswith('[1] Answer source')
    assert references[1].startswith('[2] Introduction')
    assert not any('Unused' in b.text for b in with_answers)


def test_missing_sources_remain_explicit_after_document_grouping(row):
    for source in row['sources']:
        source['document_id'] = 'same-document'
    row['asset']['sections'][0]['citation_ids'].append('declared-missing')
    row['asset']['sections'][0]['text'] += f' Missing [{UNKNOWN}]. Declared [declared-missing].'
    blocks = _material_blocks(row, True, 'en')[2]
    text = '\n'.join(b.text for b in blocks)
    assert 'Missing [?]. Declared [2].' in text
    assert '[2] Source unavailable.' in text and '[?] Source unavailable.' in text
    assert UNKNOWN not in text and 'declared-missing' not in text
